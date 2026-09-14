"""Tests for Drakopoulos source-grounded calendering mapping and fail-closed semantics."""
from __future__ import annotations

from pathlib import Path
import pytest

from src.datasets.battery_process.drakopoulos_graphite import (
    DrakopoulosGraphiteAdapter,
    SourceSemanticValidationError,
)
from src.process.contracts import ProcessStage

DATA_ROOT = Path("data/external/drakopoulos_graphite/1")
RAW_DIR = DATA_ROOT / "raw"
HAS_DATA = (RAW_DIR / "ASC-Cell_Data-Azar-Stavros.xlsx").exists() and (RAW_DIR / "ASC_results-live.xlsx").exists()


@pytest.mark.skipif(not HAS_DATA, reason="Requires Drakopoulos raw dataset")
def test_asc_calendering_source_grounding():
    """Verify calendering is mapped from ASC_results-live.xlsx and provenanced."""
    adapter = DrakopoulosGraphiteAdapter(DATA_ROOT)
    runs = adapter._parse_raw()

    asc_runs = [r for r in runs if "PROSPECTIVE" in str(r.provenance.processing_parameters.get("partition", ""))]
    assert len(asc_runs) == 108

    # In ASC dataset, Case 39 was calendered and Case 40 was uncalendered
    case_39_runs = [r for r in asc_runs if r.stages[4].provenance and r.stages[4].provenance.processing_parameters.get("source_case_id") == "Case 39"]
    case_40_runs = [r for r in asc_runs if r.stages[4].provenance and r.stages[4].provenance.processing_parameters.get("source_case_id") == "Case 40"]

    assert len(case_39_runs) >= 3
    assert len(case_40_runs) >= 3

    for r in case_39_runs:
        cal_stage = next(s for s in r.stages if s.stage_type == ProcessStage.CALENDERING)
        assert cal_stage.controls["calendering_applied"].value == 1.0
        assert cal_stage.provenance is not None
        assert cal_stage.provenance.processing_parameters.get("calendering_source") == "ASC_results-live.xlsx:Sample_Sheet_Section"
        assert cal_stage.provenance.processing_parameters.get("mapping_rule") == "EXPLICIT_SAMPLE_SHEET_CALENDARING_HEADING"

    for r in case_40_runs:
        cal_stage = next(s for s in r.stages if s.stage_type == ProcessStage.CALENDERING)
        assert cal_stage.controls["calendering_applied"].value == 0.0
        assert cal_stage.provenance is not None
        assert cal_stage.provenance.processing_parameters.get("calendering_source") == "ASC_results-live.xlsx:Sample_Sheet_Section"
        assert cal_stage.provenance.processing_parameters.get("mapping_rule") == "EXPLICIT_SAMPLE_SHEET_CALENDARING_HEADING"


@pytest.mark.skipif(not HAS_DATA, reason="Requires Drakopoulos raw dataset")
def test_as_calendering_provenance():
    """Verify AS runs have explicit calendering provenance."""
    adapter = DrakopoulosGraphiteAdapter(DATA_ROOT)
    runs = adapter._parse_raw()

    as_runs = [r for r in runs if "HISTORICAL" in str(r.provenance.processing_parameters.get("partition", ""))]
    assert len(as_runs) > 0

    for r in as_runs:
        cal_stage = next(s for s in r.stages if s.stage_type == ProcessStage.CALENDERING)
        assert cal_stage.controls["calendering_applied"].value in (0.0, 1.0)
        assert cal_stage.provenance is not None
        assert cal_stage.provenance.processing_parameters.get("calendering_source") == "AS-Cell_Data:Column_2_Calendared"
        assert cal_stage.provenance.processing_parameters.get("mapping_rule") == "NON_EMPTY_CALENDERED_CELL"


def test_fail_closed_on_missing_headers(tmp_path):
    """Verify SourceSemanticValidationError raised if required columns are missing."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FInal_All_Cell_Data"

    headers_r1 = ["Case", "Cell ID", "Active Mass", "Electrode thickness"]
    headers_r2 = ["", "", "mg", "um"]
    ws.append(headers_r1)
    ws.append(headers_r2)
    ws.append(["Case 99", "ASC-999", 10.0, 50.0])

    excel_path = tmp_path / "ASC-Cell_Data-Azar-Stavros.xlsx"
    wb.save(excel_path)
    (tmp_path / "ASC_results-live.xlsx").touch()

    adapter = DrakopoulosGraphiteAdapter(tmp_path)
    with pytest.raises(SourceSemanticValidationError, match="Required ASC columns not found"):
        adapter._parse_asc_workbook(excel_path, None)


def test_fail_closed_on_ambiguous_calendering(tmp_path):
    """Verify SourceSemanticValidationError raised if cell cannot be resolved in calendering map."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FInal_All_Cell_Data"

    headers_r1 = [
        "Case", "Cell Id", "", "", "", "", "Thickness (um)", "", "Active Mass  (mg)",
        "Cell Capacity (mAh)", "Porosity (%)", "INK", "Temperaure( C)", "Speed (m/min)",
        "Gap size(um)", "", "A%", "C%", "B1% (CMC)", "B2%", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "D30(mA.h)"
    ]
    headers_r2 = [""] * len(headers_r1)
    ws.append(headers_r1)
    ws.append(headers_r2)
    # Row with an unmapped cell
    row_data = [
        "Case 99", "ASC-999", "", "", "", "", 50.0, "", 10.0,
        3.0, 30.0, "Sample1", 60.0, 0.2,
        150.0, "", 93.0, 3.0, 2.0, 2.0, "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", 300.0
    ]
    ws.append(row_data)

    excel_path = tmp_path / "ASC-Cell_Data-Azar-Stavros.xlsx"
    wb.save(excel_path)

    # Empty ASC_results-live.xlsx so ASC-999 is definitely absent
    live_wb = openpyxl.Workbook()
    live_wb.save(tmp_path / "ASC_results-live.xlsx")

    adapter = DrakopoulosGraphiteAdapter(tmp_path)
    with pytest.raises(SourceSemanticValidationError, match="CALENDERING_STATUS_AMBIGUOUS"):
        adapter._parse_asc_workbook(excel_path, None)
